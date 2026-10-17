"""
Write VineSplit input tables from Supplementary Table 2 of Ciubotaru et al. 2021,
Front. Plant Sci. 12:693887 (doi:10.3389/fpls.2021.693887).

One file per genotype, year and time point, one row per compound:
    CAS | name | mean healthy | SE healthy | mean diseased | SE diseased
healthy = NI (mock), diseased = I (inoculated). Values copied as published, 'n.d.' kept
(the program reads it as 0). Compounds without CAS (unknowns) are left out.

Usage:  python tools/build_vineyard_profiles.py Table_2.XLSX Bianca data/vineyard
"""
import sys, os
import numpy as np
import pandas as pd
import openpyxl

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def read_table(path):
    ws = openpyxl.load_workbook(path, data_only=True)['VOCs']
    rows = []
    blocks = [(2, 3, 5, 84), (86, 87, 89, ws.max_row)]
    for hrow, yrow, r0, r1 in blocks:
        groups = []
        for c in range(4, 36, 4):
            groups.append([c, str(ws.cell(hrow, c).value).replace(' ', ''), ws.cell(yrow, c).value])
        for i, g in enumerate(groups):
            if g[2] is None:
                g[2] = groups[i - 2][2]
        geno = None
        for r in range(r0, r1 + 1):
            if ws.cell(r, 1).value:
                geno = str(ws.cell(r, 1).value).strip()
            comp = ws.cell(r, 3).value
            if not comp or comp == 'Compounds':
                continue
            for c, t, y in groups:
                rows.append(dict(genotype=geno, compound=str(comp).strip(), hpi=t, year=int(y),
                                 NI=num(ws.cell(r, c).value), NI_se=num(ws.cell(r, c + 1).value),
                                 I=num(ws.cell(r, c + 2).value), I_se=num(ws.cell(r, c + 3).value)))
    df = pd.DataFrame(rows)
    # the same Pinot Noir 2017 data appears in both blocks -> keep one copy
    return df.drop_duplicates(subset=['genotype', 'compound', 'year', 'hpi'], keep='first')


CAS = {'(E)-2-Hexenol': '928-95-0', '1-hexanol': '111-27-3', '1-hexanol-2-ethyl': '104-76-7',
       '1-Octen-3-ol': '3391-86-4', '2-hexenal': '6728-26-3', 'nonanal': '124-19-6',
       'benzaldehyde': '100-52-7', 'cis-3-hexenyl benzoate': '25152-85-6', 'methyl-salicylate': '119-36-8',
       'farnesene': '502-61-4', 'linalool': '78-70-6', 'E-nerolidol': '40716-66-3', 'neral': '106-26-3'}


def build(df, genotype):
    """Long table -> {(year, hpi): DataFrame in the input-table layout}."""
    out = {}
    g = df[df.genotype == genotype]
    for (yr, hpi), sub in g.groupby(['year', 'hpi']):
        if sub[['NI', 'I']].notna().sum().sum() == 0:
            continue                       # this year/time point was not measured for this genotype
        rows = []
        for _, r in sub.iterrows():
            if r.compound not in CAS:
                continue
            nd = lambda v: 'n.d.' if pd.isna(v) else repr(float(v))   # shortest exact text, as in the source
            rows.append({'CAS': CAS[r.compound], 'name': r.compound,
                         'mean healthy': nd(r.NI), 'SE healthy': nd(r.NI_se),
                         'mean diseased': nd(r.I), 'SE diseased': nd(r.I_se)})
        out[(int(yr), hpi)] = pd.DataFrame(rows)
    return out


if __name__ == '__main__':
    src, genotype, outdir = sys.argv[1], sys.argv[2], sys.argv[3]
    os.makedirs(outdir, exist_ok=True)
    stem = genotype.lower().replace(' ', '_')
    for (y, h), t in build(read_table(src), genotype).items():
        f = os.path.join(outdir, f'{stem}_{y}_{h}.csv')
        t.to_csv(f, index=False)
        print(f'{f}: {len(t)} compounds')
